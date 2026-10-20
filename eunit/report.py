"""Reports: JUnit XML (CI), JSON (tools) and Excel (traceability)."""
import json
import xml.etree.ElementTree as ET
from pathlib import Path

from . import EUnitError
from .checks import ERROR, FAIL, PASS, SKIP


def _ensure_dir(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _steps_text(r):
    lines = [f"{s.index}. [{s.status}] {s.description}: {s.message}" for s in r.steps]
    if r.error:
        lines.append(r.error)
    return "\n".join(lines)


def write_junit(result, path):
    meta = result.meta()
    ts = ET.Element("testsuite", name=f"{result.suite.name}.{result.runner['runner']}",
                    tests=str(len(result.tests)), failures=str(meta[FAIL]), errors=str(meta[ERROR]),
                    skipped=str(meta[SKIP]), time=f"{result.duration:.3f}", timestamp=result.started)
    props = ET.SubElement(ts, "properties")
    for k, v in {**result.runner, "hash": meta["hash"], "license": meta["license"]}.items():
        ET.SubElement(props, "property", name=str(k), value=str(v))
    for r in result.tests:
        tc = ET.SubElement(ts, "testcase", classname=f"{result.suite.name}.{result.runner['runner']}",
                           name=f"{r.test.id} {r.test.title}", time=f"{r.duration:.3f}")
        text = _steps_text(r)
        if r.status == FAIL:
            first = next(s for s in r.steps if s.status == FAIL)
            ET.SubElement(tc, "failure", message=f"{first.description}: {first.message}").text = text
        elif r.status == ERROR:
            msg = r.error or next(s.message for s in r.steps if s.status == ERROR)
            ET.SubElement(tc, "error", message=msg).text = text
        elif r.status == SKIP:
            ET.SubElement(tc, "skipped")
        ET.SubElement(tc, "system-out").text = text
    ET.indent(ts)
    ET.ElementTree(ts).write(_ensure_dir(path), encoding="utf-8", xml_declaration=True)


def to_dict(result):
    return {
        **result.meta(),
        "runner": result.runner,
        "compare": result.compare,
        "tests": [{
            "id": r.test.id, "title": r.test.title, "requirements": r.test.requirements,
            "function": r.test.call.name, "status": r.status, "samples": r.samples,
            "duration_s": round(r.duration, 4), "error": r.error,
            "cycles": None if r.cycles is None else {
                "calls": r.cycles.calls, "min": r.cycles.min, "max": r.cycles.max,
                "mean": round(r.cycles.mean, 2)},
            "steps": [{"index": s.index, "description": s.description, "status": s.status,
                       "message": s.message, "stats": s.stats} for s in r.steps],
        } for r in result.tests],
    }


def write_json(result, path):
    _ensure_dir(path).write_text(json.dumps(to_dict(result), indent=2, default=float), encoding="utf-8")


_FILL = {PASS: "C6EFCE", FAIL: "FFC7CE", ERROR: "FFC7CE", SKIP: "EDEDED"}


def write_xlsx(result, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as e:
        raise EUnitError("openpyxl is required for the Excel report: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    meta = result.meta()
    rows = [("Suite", meta["suite"]), ("Specification", meta["spec"]), ("Suite hash", meta["hash"]),
            ("Status", meta["status"]), ("Started", meta["started"]), ("Duration (s)", meta["duration_s"]),
            ("Passed", meta[PASS]), ("Failed", meta[FAIL]), ("Errors", meta[ERROR]),
            ("Skipped", meta[SKIP]), ("e2Unit", meta["eunit"]), ("License", meta["license"] or "")]
    rows += [(f"runner.{k}", str(v)) for k, v in result.runner.items()]
    if result.compare:
        rows += [(f"compare.{k}", str(v)) for k, v in result.compare.items()]
    for row in rows:
        ws.append(row)
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws["B4"].fill = PatternFill("solid", fgColor=_FILL[meta["status"]])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90

    wt = wb.create_sheet("Tests")
    header = ["Test Id", "Requirement", "Title", "Function", "Test status", "Step", "Check",
              "Step status", "Result", "Cycles max"]
    wt.append(header)
    for c in wt[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="44546A")
    for r in result.tests:
        cyc = r.cycles.max if r.cycles is not None and r.cycles.calls else None
        steps = r.steps or [None]
        for s in steps:
            wt.append([r.test.id, ", ".join(r.test.requirements), r.test.title, r.test.call.name, r.status,
                       f"{r.test.id}.{s.index}" if s else "", s.description if s else "",
                       s.status if s else "", s.message if s else (r.error or ""), cyc])
            row = wt.max_row
            wt.cell(row, 5).fill = PatternFill("solid", fgColor=_FILL[r.status])
            if s:
                wt.cell(row, 8).fill = PatternFill("solid", fgColor=_FILL[s.status])
            wt.cell(row, 9).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCDEFGHIJ", (14, 18, 36, 14, 11, 14, 40, 11, 90, 11)):
        wt.column_dimensions[col].width = width
    wt.freeze_panes = "A2"
    wt.auto_filter.ref = wt.dimensions
    wb.save(_ensure_dir(path))
